from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

from fastapi import HTTPException
from decimal import Decimal, ROUND_HALF_UP
from sqlalchemy.exc import SQLAlchemyError
from openpyxl.utils.datetime import from_excel
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.repositories.bkorder_repository import BKOrderRepository
from app.schemas.bkorder_schema import (
    BKOrderCreate,
    BKOrderImportGroup,
    BKOrderImportCommitResponse,
    BKOrderImportPreviewResponse,
    BKOrderImportRow,
    BKOrderUpdate,
)

EXCEL_REPEAT_COLUMNS = 14

IMPORT_FIELD_ALIASES = {
    "tgl": "order_date",
    "order_date": "order_date",
    "no.ord": "order_number",
    "no ord": "order_number",
    "order_number": "order_number",
    "po date": "po_date",
    "do number": "do_number",
    "deliv. date": "delivery_date",
    "pr": "customer_name",
    "langganan": "customer_name",
    "ukuran": "size",
    "jenis bahan": "material_type",
    "jenis cetak": "print_type",
    "spesifikasi": "specification",
    "sat": "unit",
    "satuan": "unit",
    "keping": "quantity",
    "rim": "rim",
    "harga": "price",
}

def _normalize_dates(values: Optional[List[Optional[str]]]) -> List[Optional[str]]:
    values = values or []
    normalized = list(values[:EXCEL_REPEAT_COLUMNS])
    while len(normalized) < EXCEL_REPEAT_COLUMNS:
        normalized.append(None)
    return normalized


def _normalize_numbers(values: Optional[List[Optional[float]]]) -> List[float]:
    values = values or []
    normalized: List[float] = []
    for value in values[:EXCEL_REPEAT_COLUMNS]:
        try:
            normalized.append(float(value or 0))
        except (TypeError, ValueError):
            normalized.append(0)
    while len(normalized) < EXCEL_REPEAT_COLUMNS:
        normalized.append(0)
    return normalized


def _normalize_optional_numbers(values: Optional[List[Optional[float]]]) -> List[Optional[float]]:
    values = values or []
    normalized: List[Optional[float]] = []
    for value in values[:EXCEL_REPEAT_COLUMNS]:
        if value in (None, ""):
            normalized.append(None)
            continue
        try:
            normalized.append(float(value))
        except (TypeError, ValueError):
            normalized.append(None)
    while len(normalized) < EXCEL_REPEAT_COLUMNS:
        normalized.append(None)
    return normalized


def _auto_total_keping(values: List[float]) -> Decimal:
    return Decimal(str(sum(values)))


def _round_money(value) -> Decimal:
    return Decimal(value or 0).quantize(Decimal("1"), rounding=ROUND_HALF_UP)


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower()


def _normalize_date_value(value: Any, workbook_epoch=None):
    if value in (None, ""):
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value, epoch=workbook_epoch).date()
        except Exception:
            return None
    if isinstance(value, str):
        value = value.strip()
        if value.lower() == "segera":
            return value
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
        return value
    if hasattr(value, "date"):
        try:
            return value.date()
        except Exception:
            return None
    return value


def _parse_excel_number(value: Any):
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return None


def _parse_excel_unit(value: Any):
    text = str(value or "").strip().upper()
    if text in {"RIM", "KEPING"}:
        return text
    return None


def _clean_optional_text(value: Any):
    if value in (None, ""):
        return None
    text = str(value).strip()
    return text or None


def _field_from_row(row: List[Any], headers: Dict[str, int], label: str):
    column = headers.get(label)
    if not column:
        return None
    return row[column - 1].value


def _row_to_data(
    row: List[Any],
    headers: Dict[str, int],
    carry: Dict[str, Any],
    workbook_epoch=None,
):
    row_warnings: List[str] = []

    def read(label: str):
        return _field_from_row(row, headers, label)

    def carry_or(value, key):
        if value in (None, ""):
            return carry.get(key)
        return value

    tgl_value = _normalize_date_value(read("tgl"), workbook_epoch)
    tgl_value = carry_or(tgl_value, "order_date")
    if isinstance(tgl_value, str) and tgl_value.lower() == "segera":
        row_warnings.append("TGL berisi teks")

    no_ord_value = _clean_optional_text(read("no.ord"))
    no_ord_value = carry_or(no_ord_value, "order_number")

    po_date_value = _normalize_date_value(read("po date"), workbook_epoch)
    po_date_value = carry_or(po_date_value, "po_date")

    do_number_value = _clean_optional_text(read("do number"))
    do_number_value = carry_or(do_number_value, "do_number")

    customer_value = _clean_optional_text(read("pr")) or _clean_optional_text(read("langganan"))
    customer_value = carry_or(customer_value, "customer_name")

    delivery_value = _normalize_date_value(read("deliv. date"), workbook_epoch)
    delivery_warning = None
    if isinstance(delivery_value, str):
        if delivery_value.lower() == "segera":
            delivery_warning = "Deliv. Date berisi Segera"
            delivery_value = None
        else:
            delivery_warning = "Deliv. Date non-date"
            delivery_value = None
    delivery_value = carry_or(delivery_value, "delivery_date")

    unit_value = _parse_excel_unit(read("sat"))
    unit_value = carry_or(unit_value, "unit")

    quantity_value = _parse_excel_number(read("keping"))
    rim_value = _parse_excel_number(read("rim"))
    price_value = _parse_excel_number(read("harga"))

    partial_delivery_values: List[Optional[str]] = []
    for col_idx in range(15, 35):
        cell_value = row[col_idx - 1].value if len(row) >= col_idx else None
        parsed = _normalize_date_value(cell_value, workbook_epoch)
        if isinstance(parsed, date):
            partial_delivery_values.append(parsed.isoformat())
        elif isinstance(parsed, datetime):
            partial_delivery_values.append(parsed.date().isoformat())
        elif isinstance(parsed, str) and parsed:
            partial_delivery_values.append(parsed)
        else:
            partial_delivery_values.append(None)

    carry_updates = {
        "order_date": tgl_value,
        "order_number": no_ord_value,
        "po_date": po_date_value,
        "do_number": do_number_value,
        "customer_name": customer_value,
        "delivery_date": delivery_value,
        "unit": unit_value,
    }

    data = {
        "order_date": tgl_value,
        "order_number": no_ord_value,
        "po_date": po_date_value,
        "do_number": do_number_value,
        "delivery_date": delivery_value,
        "customer_name": customer_value,
        "size": _clean_optional_text(read("ukuran")),
        "material_type": _clean_optional_text(read("jenis bahan")),
        "print_type": _clean_optional_text(read("jenis cetak")),
        "specification": _clean_optional_text(read("spesifikasi")),
        "unit": unit_value,
        "quantity": quantity_value,
        "rim": rim_value,
        "price": price_value,
        "status": "OPEN",
        "notes": None,
        "delivery_completed_dates": partial_delivery_values,
        "partial_billing_quantities": [0] * EXCEL_REPEAT_COLUMNS,
        "partial_billing_input_quantities": [None] * EXCEL_REPEAT_COLUMNS,
        "warnings": [warning for warning in [delivery_warning, *row_warnings] if warning],
    }

    return data, carry_updates


def _group_key_for_order(data: Dict[str, Any], fallback_key: str) -> str:
    order_number = str(data.get("order_number") or "").strip()
    return order_number or fallback_key


def _serialize_date(value: Any):
    if isinstance(value, date):
        return value.isoformat()
    return value


def _row_error(row_number: int, errors: List[str], data: Dict[str, Any], is_valid: bool):
    return BKOrderImportRow(row_number=row_number, errors=errors, data=data, is_valid=is_valid)


class BKOrderService:
    def __init__(self, db: Session):
        self.repository = BKOrderRepository(db)

    def _prepare_create(self, data: BKOrderCreate) -> BKOrderCreate:
        payload = data.model_dump()
        payload["delivery_completed_dates"] = _normalize_dates(payload.get("delivery_completed_dates"))
        payload["partial_billing_quantities"] = _normalize_numbers(payload.get("partial_billing_quantities"))
        payload["partial_billing_input_quantities"] = _normalize_optional_numbers(
            payload.get("partial_billing_input_quantities")
        )

        if payload.get("total_keping") in (None, 0, Decimal("0"), ""):
            payload["total_keping"] = _auto_total_keping(payload["partial_billing_quantities"])

        if "price" in payload:
            payload["price"] = _round_money(payload.get("price"))

        return BKOrderCreate(**payload)

    def _prepare_update(self, data: BKOrderUpdate) -> BKOrderUpdate:
        payload = data.model_dump(exclude_unset=True)

        if "delivery_completed_dates" in payload:
            payload["delivery_completed_dates"] = _normalize_dates(payload.get("delivery_completed_dates"))

        if "partial_billing_quantities" in payload:
            payload["partial_billing_quantities"] = _normalize_numbers(payload.get("partial_billing_quantities"))
            if payload.get("total_keping") in (None, 0, Decimal("0"), ""):
                payload["total_keping"] = _auto_total_keping(payload["partial_billing_quantities"])

        if "partial_billing_input_quantities" in payload:
            payload["partial_billing_input_quantities"] = _normalize_optional_numbers(
                payload.get("partial_billing_input_quantities")
            )

        if "price" in payload:
            payload["price"] = _round_money(payload.get("price"))

        return BKOrderUpdate(**payload)

    def get_all(
        self,
        search: Optional[str] = None,
        status: Optional[str] = None,
        month: Optional[int] = None,
        year: Optional[int] = None,
    ):
        return self.repository.get_all(search=search, status=status, month=month, year=year)

    def get_by_id(self, bkorder_id: int):
        obj = self.repository.get_by_id(bkorder_id)
        if not obj:
            raise HTTPException(status_code=404, detail="BKOrder tidak ditemukan")
        return obj

    def create(self, data: BKOrderCreate):
        return self.repository.create(self._prepare_create(data))

    def update(self, bkorder_id: int, data: BKOrderUpdate):
        obj = self.repository.update(bkorder_id, self._prepare_update(data))
        if not obj:
            raise HTTPException(status_code=404, detail="BKOrder tidak ditemukan")
        return obj

    def delete(self, bkorder_id: int):
        obj = self.repository.delete(bkorder_id)
        if not obj:
            raise HTTPException(status_code=404, detail="BKOrder tidak ditemukan")
        return {"message": "BKOrder berhasil dihapus"}

    def timeline(self, bkorder_id: int):
        obj = self.get_by_id(bkorder_id)
        kirim_dates = obj.delivery_completed_dates or []
        kirim_dates = [value for value in kirim_dates if value]

        timeline = [
            {"label": "TGL", "date": obj.order_date, "description": "Tanggal order masuk"},
            {"label": "PO Date", "date": obj.po_date, "description": "Tanggal PO"},
            {"label": "Deliv. Date", "date": obj.delivery_date, "description": "Tanggal rencana delivery"},
        ]

        for index, value in enumerate(kirim_dates, start=1):
            timeline.append(
                {
                    "label": f"TGL KIRIM / SELESAI {index}",
                    "date": value,
                    "description": "Tanggal kirim / selesai sesuai kolom Excel",
                }
            )

        return timeline

    def preview_import(self, file_bytes: bytes, filename: str) -> BKOrderImportPreviewResponse:
        if not filename.lower().endswith(".xlsx"):
            raise HTTPException(status_code=400, detail="Saat ini import hanya mendukung file .xlsx")

        workbook = load_workbook(BytesIO(file_bytes), data_only=True)
        sheet = workbook.worksheets[0]
        headers: Dict[str, int] = {}
        rows: List[BKOrderImportRow] = []
        groups: List[BKOrderImportGroup] = []
        valid_count = 0
        error_count = 0
        duplicate_count = 0
        carry: Dict[str, Any] = {}
        current_group_key: Optional[str] = None
        current_group: Optional[BKOrderImportGroup] = None

        for cell in sheet[1]:
            key = _normalize_header(cell.value)
            if key:
                headers[key] = cell.column

        for excel_row_number, row in enumerate(sheet.iter_rows(min_row=3), start=3):
            if all(cell.value in (None, "") for cell in row):
                continue

            data, carry = _row_to_data(list(row), headers, carry, workbook.epoch)
            order_number = str(data.get("order_number") or "").strip()
            order_date = data.get("order_date")
            unit = data.get("unit")
            row_errors: List[str] = []
            row_warnings = list(data.get("warnings") or [])

            if not order_number:
                row_errors.append("NO.ORD wajib diisi")
            if not order_date or isinstance(order_date, str):
                row_errors.append("TGL harus valid")
            if unit and unit not in {"RIM", "KEPING"}:
                row_errors.append("SAT hanya boleh RIM atau KEPING")
            if data.get("delivery_date") == "Segera":
                row_warnings.append("Deliv. Date berisi Segera")

            db_duplicate = False
            if order_number and self.repository.get_by_order_number(order_number):
                db_duplicate = True

            if db_duplicate:
                row_warnings.append("NO.ORD sudah ada di database")

            is_valid = len(row_errors) == 0 and not db_duplicate
            is_header_row = bool(
                data.get("order_date")
                or data.get("order_number")
                or data.get("po_date")
                or data.get("do_number")
                or data.get("customer_name")
            )
            if is_valid:
                valid_count += 1
            else:
                error_count += 1
            if db_duplicate:
                duplicate_count += 1

            group_key = _group_key_for_order(data, f"row-{excel_row_number}")
            if current_group_key != group_key:
                current_group = BKOrderImportGroup(
                    order_number=order_number or None,
                    customer_name=data.get("customer_name"),
                    order_date=_serialize_date(data.get("order_date")),
                    po_date=_serialize_date(data.get("po_date")),
                    do_number=data.get("do_number"),
                    delivery_date=_serialize_date(data.get("delivery_date")),
                )
                groups.append(current_group)
                current_group_key = group_key

            preview_row = BKOrderImportRow(
                row_number=excel_row_number,
                data={
                    "order_date": _serialize_date(order_date),
                    "order_number": order_number or None,
                    "po_date": _serialize_date(data.get("po_date")),
                    "do_number": data.get("do_number"),
                    "delivery_date": _serialize_date(data.get("delivery_date")),
                    "customer_name": data.get("customer_name"),
                    "size": data.get("size"),
                    "material_type": data.get("material_type"),
                    "print_type": data.get("print_type"),
                    "specification": data.get("specification"),
                    "unit": data.get("unit"),
                    "quantity": data.get("quantity"),
                    "rim": data.get("rim"),
                    "price": data.get("price"),
                    "delivery_completed_dates": data.get("delivery_completed_dates"),
                    "partial_billing_quantities": data.get("partial_billing_quantities"),
                    "partial_billing_input_quantities": data.get("partial_billing_input_quantities"),
                    "status": data.get("status"),
                    "notes": data.get("notes"),
                },
                errors=row_errors,
                warnings=row_warnings,
                is_valid=is_valid,
                is_header_row=is_header_row,
            )
            rows.append(preview_row)
            if current_group is not None:
                current_group.rows.append(preview_row)
                current_group.row_count += 1
                if is_valid:
                    current_group.valid_count += 1
                else:
                    current_group.error_count += 1
                if db_duplicate:
                    current_group.duplicate_count += 1

        return BKOrderImportPreviewResponse(
            groups=groups,
            rows=rows,
            valid_count=valid_count,
            error_count=error_count,
            duplicate_count=duplicate_count,
        )

    def commit_import(self, rows: List[dict[str, Any]]) -> BKOrderImportCommitResponse:
        success_count = 0
        failed_count = 0
        duplicate_count = 0
        errors: List[dict[str, Any]] = []

        for index, row in enumerate(rows, start=1):
            row_number = int(row.get("row_number") or index)
            data = row.get("data") or {}
            row_errors = list(row.get("errors") or [])
            order_number = str(data.get("order_number") or "").strip()
            order_date = _normalize_date_value(data.get("order_date"))
            po_date = _normalize_date_value(data.get("po_date"))
            delivery_date = _normalize_date_value(data.get("delivery_date"))

            if not order_number:
                row_errors.append("NO.ORD wajib diisi")
            if not order_date or isinstance(order_date, str):
                row_errors.append("TGL harus valid")
            if data.get("unit") and str(data.get("unit")).upper() not in {"RIM", "KEPING"}:
                row_errors.append("SAT hanya boleh RIM atau KEPING")

            if row_errors:
                failed_count += 1
                errors.append({"row": row_number, "message": "; ".join(row_errors)})
                continue

            payload = {
                "order_date": order_date,
                "order_number": order_number,
                "po_date": po_date,
                "do_number": data.get("do_number"),
                "delivery_date": delivery_date if not isinstance(delivery_date, str) else None,
                "customer_name": data.get("customer_name"),
                "size": data.get("size"),
                "material_type": data.get("material_type"),
                "print_type": data.get("print_type"),
                "specification": data.get("specification"),
                "unit": str(data.get("unit")).upper() if data.get("unit") else None,
                "quantity": data.get("quantity") or 0,
                "rim": data.get("rim") or 0,
                "price": data.get("price") or 0,
                "delivery_completed_dates": data.get("delivery_completed_dates") or [None] * EXCEL_REPEAT_COLUMNS,
                "partial_billing_quantities": data.get("partial_billing_quantities") or [0] * EXCEL_REPEAT_COLUMNS,
                "partial_billing_input_quantities": data.get("partial_billing_input_quantities") or [None] * EXCEL_REPEAT_COLUMNS,
                "total_keping": data.get("quantity") or 0,
                "status": data.get("status") or "OPEN",
                "notes": data.get("notes"),
            }
            try:
                self.repository.create(self._prepare_create(BKOrderCreate(**payload)))
                success_count += 1
            except SQLAlchemyError as exc:
                self.repository.db.rollback()
                failed_count += 1
                duplicate_count += 1 if "duplicate" in str(exc).lower() else 0
                errors.append({"row": row_number, "message": str(exc).split("\n")[0]})
            except Exception as exc:
                self.repository.db.rollback()
                failed_count += 1
                errors.append({"row": row_number, "message": str(exc)})

        return BKOrderImportCommitResponse(
            success=failed_count == 0,
            success_count=success_count,
            failed_count=failed_count,
            duplicate_count=duplicate_count,
            errors=errors,
        )
